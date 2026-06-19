"""Regression tests for Korean / UTF-8 handling in file flows.

These exist to keep three classes of subtle bugs from coming back:
  1. PDF rendered with a non-CJK font → Korean glyphs become squares
  2. Content-Disposition with raw non-ASCII filename → browser saves "_____.csv"
  3. CSV without a UTF-8 BOM → Excel opens it as cp949 and shows mojibake
"""
from decimal import Decimal


# --- PDF Korean rendering ---------------------------------------------------


def test_pdf_renders_korean_text():
    import pytest
    from io import BytesIO

    pypdf = pytest.importorskip("pypdf", reason="pypdf는 선택적 dev 의존성")
    PdfReader = pypdf.PdfReader

    from app.core.exports import render_pdf_bytes

    pdf_bytes = render_pdf_bytes(
        rows=[["웰그린 라들러 레몬", 100], ["바이젠 맥주", 50]],
        headers=["품목명", "재고"],
        title="월간 보고서",
    )
    text = PdfReader(BytesIO(pdf_bytes)).pages[0].extract_text()
    # Hangul Syllables: U+AC00 .. U+D7A3
    korean = [c for c in text if 0xAC00 <= ord(c) <= 0xD7A3]
    assert len(korean) >= 10, f"Korean glyphs missing — extracted: {text!r}"


def test_pdf_filename_uses_rfc5987():
    from app.core.exports import _content_disposition

    h = _content_disposition("직원목록.csv")
    # Must include the UTF-8 form so browsers preserve the Korean name.
    assert "filename*=UTF-8''" in h
    assert "%EC%A7%81" in h  # 직 (the first char) percent-encoded
    # Also include an ASCII fallback for ancient clients
    assert 'filename="' in h


# --- CSV Korean export ------------------------------------------------------


def test_csv_export_has_utf8_bom_and_korean(client, admin_auth, db_session):
    from app.modules.hr.models import Employee

    db_session.add(
        Employee(
            employee_no="TEST-001",
            full_name="김한글",
            email="kim@test.com",
            position="과장",
            salary=Decimal("0"),
        )
    )
    db_session.commit()

    res = client.get(
        "/api/hr/employees/export?format=csv", headers=admin_auth["headers"]
    )
    assert res.status_code == 200
    assert "text/csv" in res.headers["content-type"]
    assert "charset=utf-8" in res.headers["content-type"]
    body = res.content
    # UTF-8 BOM so Excel auto-detects encoding.
    assert body.startswith(b"\xef\xbb\xbf"), "CSV must start with UTF-8 BOM"
    # Korean payload survives end-to-end.
    decoded = body.decode("utf-8-sig")
    assert "김한글" in decoded
    assert "과장" in decoded
    # RFC 5987 filename
    cd = res.headers["content-disposition"]
    assert "filename*=UTF-8''" in cd


# --- XLSX Korean export -----------------------------------------------------


def test_xlsx_export_keeps_korean(client, admin_auth, db_session):
    from io import BytesIO

    from openpyxl import load_workbook

    from app.modules.hr.models import Employee

    db_session.add(
        Employee(
            employee_no="TEST-002",
            full_name="이서연",
            email="lee@test.com",
            position="대리",
            salary=Decimal("0"),
        )
    )
    db_session.commit()

    res = client.get(
        "/api/hr/employees/export?format=xlsx", headers=admin_auth["headers"]
    )
    assert res.status_code == 200
    wb = load_workbook(BytesIO(res.content), read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    headers_row = rows[0]
    assert "이름" in headers_row, f"Korean header missing: {headers_row}"
    body_cells = [c for r in rows[1:] for c in r if c]
    assert any("이서연" in str(c) for c in body_cells)
    assert "filename*=UTF-8''" in res.headers["content-disposition"]


# --- PDF export end-to-end (Korean filename + Korean body) ------------------


def test_pdf_export_endpoint_with_korean_name(client, admin_auth, db_session):
    from app.modules.hr.models import Employee

    db_session.add(
        Employee(
            employee_no="TEST-003",
            full_name="박도현",
            email="park@test.com",
            position="물류 매니저",
            salary=Decimal("0"),
        )
    )
    db_session.commit()

    res = client.get(
        "/api/hr/employees/export?format=pdf", headers=admin_auth["headers"]
    )
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/pdf"
    # PDF magic header
    assert res.content[:4] == b"%PDF"
    # Filename header is RFC 5987
    assert "filename*=UTF-8''" in res.headers["content-disposition"]


# --- Attachment upload with Korean filename ---------------------------------


def test_attachment_upload_preserves_korean_filename(client, admin_auth):
    """Uploaded file's display filename must round-trip Korean intact."""
    files = {"file": ("회의록_2026.pdf", b"%PDF-1.4\n%fake", "application/pdf")}
    res = client.post(
        "/api/attachments",
        headers=admin_auth["headers"],
        files=files,
    )
    assert res.status_code == 200
    body = res.json()
    assert body["filename"] == "회의록_2026.pdf"

    # Download responds with the Korean filename properly encoded.
    download = client.get(
        f"/api/attachments/{body['id']}/download",
        headers=admin_auth["headers"],
    )
    assert download.status_code == 200
    cd = download.headers.get("content-disposition", "")
    # Starlette uses lowercase utf-8 in filename*=
    assert "filename*=utf-8''" in cd.lower()
